"""Solve Task 1: the 20-day preform lot-sizing problem.

The model decides:
1. how many preforms to produce each day;
2. whether a production batch is started each day;
3. the finished-preform inventory at the end of each day; and
4. the backorders at the end of each day.

PuLP is used to build the optimisation model, and its CBC solver finds the
minimum-cost production plan.
"""

from pathlib import Path

from openpyxl import load_workbook
from pulp import (
    LpBinary,
    LpInteger,
    LpMinimize,
    LpProblem,
    LpStatus,
    LpVariable,
    PULP_CBC_CMD,
    lpSum,
    value,
)


# The Excel file is in the same folder as this Python file.
DATA_FILE = Path(__file__).resolve().parent / "ZORA_synthetic_demand_data.xlsx"
SHEET_NAME = "Confirmed_20_Day_Demand"

# Planning information from the assignment.
NUMBER_OF_DAYS = 20
INITIAL_INVENTORY = 20
MAX_BATCH_SIZE = 100

# Costs in dollars.
SETUP_COST = 1_000
VARIABLE_COST = 10
HOLDING_COST = 15
BACKORDER_COST = 100


def load_demand():
    """Read the confirmed 20-day demand from the Excel workbook."""
    workbook = load_workbook(DATA_FILE, data_only=True, read_only=True)
    worksheet = workbook[SHEET_NAME]

    # Find the column called "confirmed_demand_preforms" in the first row.
    headings = [cell.value for cell in worksheet[1]]
    demand_column = headings.index("confirmed_demand_preforms") + 1

    # Store demand as {day: demand}; for example, demand[1] is Day 1 demand.
    demand = {}
    for day in range(1, NUMBER_OF_DAYS + 1):
        excel_row = day + 1  # Row 1 contains the headings.
        demand[day] = int(worksheet.cell(excel_row, demand_column).value)

    workbook.close()
    return demand


def solve_lot_sizing():
    """Build and solve the mixed-integer lot-sizing model."""
    demand = load_demand()
    days = range(1, NUMBER_OF_DAYS + 1)

    # Create a minimisation model.
    model = LpProblem("Preform_Lot_Sizing", LpMinimize)

    # Decision variables for each day.
    # All quantities are whole preforms, so they are integer variables.
    production = LpVariable.dicts("Production", days, lowBound=0, cat=LpInteger)
    setup = LpVariable.dicts("Setup", days, cat=LpBinary)
    inventory = LpVariable.dicts("Inventory", days, lowBound=0, cat=LpInteger)
    backorders = LpVariable.dicts("Backorders", days, lowBound=0, cat=LpInteger)

    # Objective: minimise setup, production, inventory, and backorder costs.
    model += lpSum(
        SETUP_COST * setup[day]
        + VARIABLE_COST * production[day]
        + HOLDING_COST * inventory[day]
        + BACKORDER_COST * backorders[day]
        for day in days
    )

    for day in days:
        # No production is possible without a setup. If setup = 1, production
        # can be any whole number from 0 to the maximum batch size of 100.
        model += (
            production[day] <= MAX_BATCH_SIZE * setup[day],
            f"Batch_size_limit_day_{day}",
        )

        # Inventory balance:
        # ending net inventory = starting net inventory + production - demand.
        # Net inventory means inventory minus backorders.
        if day == 1:
            model += (
                inventory[day] - backorders[day]
                == INITIAL_INVENTORY + production[day] - demand[day],
                "Inventory_balance_day_1",
            )
        else:
            model += (
                inventory[day] - backorders[day]
                == inventory[day - 1]
                - backorders[day - 1]
                + production[day]
                - demand[day],
                f"Inventory_balance_day_{day}",
            )

    # All customer orders must be satisfied by the end of Day 20.
    model += backorders[NUMBER_OF_DAYS] == 0, "No_final_backorders"

    # Solve the model without displaying the solver's technical messages.
    model.solve(PULP_CBC_CMD(msg=False))

    if LpStatus[model.status] != "Optimal":
        raise RuntimeError(
            f"The solver did not find an optimal solution: {LpStatus[model.status]}"
        )

    # Save the optimal daily decisions in a simple list of dictionaries.
    schedule = []
    for day in days:
        schedule.append(
            {
                "day": day,
                "demand": demand[day],
                "production": round(value(production[day])),
                "setup": round(value(setup[day])),
                "inventory": round(value(inventory[day])),
                "backorders": round(value(backorders[day])),
            }
        )

    # Calculate each part of the total cost for the report.
    setup_cost_total = sum(SETUP_COST * row["setup"] for row in schedule)
    variable_cost_total = sum(VARIABLE_COST * row["production"] for row in schedule)
    holding_cost_total = sum(HOLDING_COST * row["inventory"] for row in schedule)
    backorder_cost_total = sum(BACKORDER_COST * row["backorders"] for row in schedule)

    cost_breakdown = {
        "Setup cost": setup_cost_total,
        "Variable production cost": variable_cost_total,
        "Finished-preform holding cost": holding_cost_total,
        "Backorder cost": backorder_cost_total,
        "Total cost": round(value(model.objective)),
    }

    return schedule, cost_breakdown


def print_results(schedule, cost_breakdown):
    """Print the optimal plan and cost breakdown in readable tables."""
    print("Optimal daily production plan")
    print("-" * 72)
    print(
        f"{'Day':>3} {'Demand':>8} {'Production':>12} "
        f"{'Setup':>7} {'Inventory':>11} {'Backorders':>12}"
    )
    print("-" * 72)

    for row in schedule:
        print(
            f"{row['day']:>3} {row['demand']:>8} {row['production']:>12} "
            f"{row['setup']:>7} {row['inventory']:>11} {row['backorders']:>12}"
        )

    print("\nCost breakdown")
    print("-" * 45)
    for cost_name, amount in cost_breakdown.items():
        print(f"{cost_name:<34} ${amount:>9,.2f}")


if __name__ == "__main__":
    optimal_schedule, optimal_costs = solve_lot_sizing()
    print_results(optimal_schedule, optimal_costs)
